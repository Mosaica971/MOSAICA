"""The fill-in Excel workbook of a data catalogue, and reading it back.

One sheet per table, in the style of a carbon-assessment workbook: a *Read me* sheet, a units
and an allowed-values row, drop-down lists taken from the referentials, cells to fill in yellow,
headers and computed cells locked in grey, an example row, and a *Sources* sheet. The schema is
core/data/catalogue.py; this module only lays it out and reads it back (read_workbook), so that
validation stays in one place.

Every table sheet starts with the same five rows -- title, description, provenance, value rule,
header -- so a reader never has to guess where a table starts. Records tables add three rows of
field metadata and an example row before the data.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, NamedStyle, PatternFill, Protection
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles.cell_style import StyleArray
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core.data.catalogue import Catalogue, TableSpec, excel_formula, rule_text

HEADER_ROW = 5
RECORDS_META = ("description", "unit", "allowed")  # rows 6, 7, 8 of a records sheet
EXAMPLE_ROW = 9
EXAMPLE_KEY = "EXAMPLE"
RECORDS_FIRST_ROW = 10
FIRST_ROW = 6  # set, mapping, matrix
MATRIX_META_COLUMNS = ("unit", "allowed")  # matrix sheets whose rows are variables
README, SOURCES, LISTS = "Read me", "Sources", "Lists"

INPUT = PatternFill("solid", fgColor="FFF2CC")  # yellow: to fill in
LOCKED = PatternFill("solid", fgColor="D9D9D9")  # grey: header, computed, fixed keys
EXAMPLE = PatternFill("solid", fgColor="F2F2F2")
_GREY_ITALIC = Font(italic=True, color="808080")
_TOP_WRAP = Alignment(wrap_text=True, vertical="top")

# Every cell gets one named style: one assignment per cell, and a write-only workbook (rows
# streamed to disk) -- an in-memory one holds ~1 kB per cell and a 25 000-plot sheet has
# 600 000 of them, which is where a full workbook stopped fitting in memory.
_STYLES = {
    "input": dict(fill=INPUT, protection=Protection(locked=False)),
    "locked": dict(fill=LOCKED),
    "header": dict(fill=LOCKED, font=Font(bold=True)),
    "meta": dict(fill=LOCKED, alignment=_TOP_WRAP),
    "example": dict(fill=EXAMPLE, font=_GREY_ITALIC),
    "note": dict(font=_GREY_ITALIC),
    "title": dict(font=Font(bold=True, size=14)),
    "bold": dict(font=Font(bold=True)),
    "wrap": dict(alignment=_TOP_WRAP),
}


# --- building ----------------------------------------------------------------------------------

def build_workbook(
    catalogue: Catalogue,
    sheets: list[str],
    frames: dict[str, pd.DataFrame],
    referentials: dict[str, list[str]],
    path: Path,
    spare_rows: int = 50,
) -> Path:
    """Write the workbook for `sheets` (catalogue sheet names), prefilled from `frames`.

    A sheet absent from `frames` is written empty, with `spare_rows` input rows.
    """
    wb = Workbook(write_only=True)
    for name, style in _STYLES.items():
        wb.add_named_style(NamedStyle(name=name, **style))
    lists = _Lists(referentials)
    _write_readme(_Sheet(wb, README), catalogue, sheets)
    for sheet in sheets:
        spec = catalogue.tables[sheet]
        out = _Sheet(wb, sheet)
        writer = {"set": _write_set, "mapping": _write_mapping, "records": _write_records,
                  "matrix": _write_matrix}[spec.layout]
        writer(out, spec, frames.get(sheet), lists, spare_rows)
        out.protect()
    _write_sources(_Sheet(wb, SOURCES), catalogue, sheets)
    lists.write(_Sheet(wb, LISTS))
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


class _Sheet:
    """A write-only worksheet: rows are appended in order, each cell as (value, style)."""

    def __init__(self, wb, title: str):
        self.ws = wb.create_sheet(title)
        self.row = 0
        self._styles: dict[str, StyleArray] = {}

    def _style(self, name: str) -> StyleArray:
        # Resolving a named style costs more than writing the cell: resolve once per sheet,
        # then copy the resulting style array.
        if name not in self._styles:
            prototype = WriteOnlyCell(self.ws)
            prototype.style = name
            self._styles[name] = prototype._style
        return self._styles[name]

    def append(self, *cells) -> int:
        """Append one row; each cell is a value or a (value, style) pair. Returns the row number."""
        row = []
        for cell in cells:
            value, style = cell if isinstance(cell, tuple) else (cell, None)
            if style is None:
                row.append(value)
            else:
                written = WriteOnlyCell(self.ws, value=value)
                written._style = StyleArray(self._style(style))
                row.append(written)
        self.ws.append(row)
        self.row += 1
        return self.row

    def width(self, column: int, width: float) -> None:
        self.ws.column_dimensions[get_column_letter(column)].width = width

    def validate(self, spec: dict, cells: str, lists: "_Lists") -> None:
        dv = _validation(spec, lists)
        if dv is not None:
            dv.add(cells)
            self.ws.data_validations.append(dv)

    def protect(self) -> None:
        """Lock headers and computed cells; no password -- this prevents accidents, not access."""
        protection = self.ws.protection
        protection.sheet = True
        protection.formatColumns = False
        protection.formatRows = False
        protection.autoFilter = False
        protection.sort = False


class _Lists:
    """The hidden sheet holding the drop-down lists, one column per referential used."""

    def __init__(self, referentials: dict[str, list[str]]):
        self.referentials = referentials
        self.used: list[str] = []

    def range_for(self, name: str) -> str | None:
        values = self.referentials.get(name)
        if not values:
            return None
        if name not in self.used:
            self.used.append(name)
        column = get_column_letter(self.used.index(name) + 1)
        return f"'{LISTS}'!${column}$2:${column}${len(values) + 1}"

    def write(self, sheet: "_Sheet") -> None:
        columns = [self.referentials[name] for name in self.used]
        sheet.append(*self.used)
        for i in range(max((len(c) for c in columns), default=0)):
            sheet.append(*(c[i] if i < len(c) else None for c in columns))
        sheet.ws.sheet_state = "hidden"


def _validation(spec: dict, lists: _Lists) -> DataValidation | None:
    """The data validation of a field spec, if it has one."""
    kind = spec.get("type", "float")
    dv = None
    if kind == "binary":
        dv = DataValidation(type="list", formula1='"0,1"')
    elif kind == "code":
        codes = ",".join(str(k) for k in spec["values"])
        dv = DataValidation(type="list", formula1=f'"{codes}"') if len(codes) < 250 else None
    elif kind == "key":
        source = lists.range_for(spec.get("referential", ""))
        dv = DataValidation(type="list", formula1=source) if source else None
    elif kind in ("int", "float"):
        dv_type = "whole" if kind == "int" else "decimal"
        low, high = spec.get("min"), spec.get("max")
        if low is not None and high is not None:
            dv = DataValidation(type=dv_type, operator="between", formula1=str(low), formula2=str(high))
        elif low is not None:
            dv = DataValidation(type=dv_type, operator="greaterThanOrEqual", formula1=str(low))
        elif high is not None:
            dv = DataValidation(type=dv_type, operator="lessThanOrEqual", formula1=str(high))
    if dv is None:
        return None
    message = rule_text(spec)[:255]
    dv.allow_blank = True
    dv.error, dv.errorTitle, dv.prompt = message, "Value not allowed", message
    dv.showErrorMessage = dv.showInputMessage = True
    return dv


def _key_spec(axis: dict) -> dict:
    return {"type": "key", "referential": axis["referential"]} if axis.get("referential") else {}


def _head(out: _Sheet, spec: TableSpec, rule: str) -> None:
    """Rows 1-4 of every table sheet: title, description, provenance, value rule."""
    out.append((spec.entry.get("title", spec.sheet), "title"))
    out.append(" ".join(str(spec.entry.get("description", "")).split()))
    out.append((f"Access level {spec.level} -- file {spec.file} -- source: "
                f"{spec.entry.get('source', 'to confirm')}", "note"))
    out.append((rule, "note"))


def _write_set(out, spec, frame, lists, spare_rows):
    _head(out, spec, "one id per row")
    out.append(("id", "header"))
    values = [] if frame is None else list(frame["id"])
    for value in values + [None] * spare_rows:
        out.append((value, "input"))
    out.width(1, 24)
    out.ws.freeze_panes = f"A{FIRST_ROW}"


def _write_mapping(out, spec, frame, lists, spare_rows):
    parent, child = spec.key_names
    _head(out, spec, f"one {parent} -> {child} link per row")
    out.append((parent, "header"), (child, "header"))
    pairs = [] if frame is None else list(zip(frame[parent], frame[child]))
    for p, c in pairs + [(None, None)] * spare_rows:
        out.append((p, "input"), (c, "input"))
    last = out.row
    for column, role in ((1, "parent"), (2, "child")):
        letter = get_column_letter(column)
        out.validate(_key_spec(spec.entry[role]), f"{letter}{FIRST_ROW}:{letter}{last}", lists)
        out.width(column, 16)
    out.ws.freeze_panes = f"A{FIRST_ROW}"


def _example_value(spec: dict, lists: _Lists):
    kind = spec.get("type", "float")
    if kind == "code":
        return int(next(iter(spec["values"])))
    if kind == "binary":
        return 0
    if kind == "key":
        values = lists.referentials.get(spec.get("referential", ""))
        return values[0] if values else None
    return spec.get("min", spec.get("max", 0))


def _write_records(out, spec, frame, lists, spare_rows):
    fields = spec.fields
    names = list(fields)
    key_name = spec.rows.get("name", "id")
    _head(out, spec, f"one row per {key_name}; rows 6-8 describe each column")
    out.append((key_name, "header"), *((name, "header") for name in names))

    def description(field_spec):
        text = field_spec.get("description", "")
        if field_spec.get("note"):
            text += " -- " + " ".join(str(field_spec["note"]).split())
        return text

    meta = {"description": description, "unit": lambda f: f.get("unit", ""), "allowed": rule_text}
    for label in RECORDS_META:
        out.append((label, "header"), *((meta[label](fields[n]), "meta") for n in names))

    example = (frame.reindex(columns=names).iloc[0].tolist() if frame is not None and len(frame)
               else [_example_value(fields[n], lists) for n in names])
    out.append((EXAMPLE_KEY, "example"), *((_cell_value(v), "example") for v in example))

    column_of = {name: get_column_letter(i + 2) for i, name in enumerate(names)}
    kinds = ["computed" if "computed" in fields[n] else "derived" if "derived" in fields[n]
             else "input" for n in names]
    keys = [] if frame is None else list(frame.index)
    records = [] if frame is None else frame.reindex(columns=names).to_numpy(dtype=object).tolist()
    for index in range(len(keys) + spare_rows):
        row = RECORDS_FIRST_ROW + index
        record = records[index] if index < len(keys) else None
        cells = [(keys[index] if record is not None else None, "input")]
        for i, (name, kind) in enumerate(zip(names, kinds)):
            if kind == "computed" and record is None:  # spare row: Excel computes it
                formula = excel_formula(fields[name]["computed"],
                                        {n: f"{column_of[n]}{row}" for n in names})
                cells.append((f'=IF($A{row}="","",{formula[1:]})', "locked"))
            elif kind == "computed":
                # A stored value is kept as is: recomputing it from rounded inputs moves it by
                # float noise (the validator accepts it within the catalogue's tolerance).
                cells.append((_cell_value(record[i]), "locked"))
            else:
                value = _cell_value(record[i]) if record is not None else None
                cells.append((value, "locked" if kind == "derived" else "input"))
        out.append(*cells)
    last = out.row
    out.validate(_key_spec(spec.rows), f"A{RECORDS_FIRST_ROW}:A{last}", lists)
    for name, kind in zip(names, kinds):
        if kind == "input":
            letter = column_of[name]
            out.validate(fields[name], f"{letter}{RECORDS_FIRST_ROW}:{letter}{last}", lists)
    out.width(1, 16)
    for i, name in enumerate(names, start=2):
        out.width(i, max(12, min(28, len(name) + 4)))
    out.ws.freeze_panes = f"B{RECORDS_FIRST_ROW}"


def _write_matrix(out, spec, frame, lists, spare_rows):
    row_name = spec.rows.get("name", "row")
    col_name = spec.columns.get("name", "column")
    row_fields = spec.row_fields
    if row_fields:
        rule = f"one row per {row_name}, each with its own unit and rule; one column per {col_name}"
    else:
        value = spec.value
        rule = (f"{row_name} x {col_name}; every value: {value.get('description', '')}"
                f"{' (' + value['unit'] + ')' if value.get('unit') else ''}; {rule_text(value)}")
    _head(out, spec, rule)
    if frame is not None:
        columns = [str(c) for c in frame.columns]
    else:
        columns = list(lists.referentials.get(spec.columns.get("referential", ""), []))
    meta = list(MATRIX_META_COLUMNS) if row_fields else []
    offset = 2 + len(meta)
    out.append((row_name, "header"), *((c, "header") for c in meta + columns))

    if row_fields:
        keys = list(row_fields)
    elif frame is not None:
        keys = list(frame.index)
    else:
        keys = list(lists.referentials.get(spec.rows.get("referential", ""), []))
    grid = {} if frame is None else dict(zip(
        frame.index, frame.reindex(columns=columns).to_numpy(dtype=object).tolist()))
    first_col = get_column_letter(offset)
    last_col = get_column_letter(offset + max(len(columns), 1) - 1)
    extra = 0 if row_fields else spare_rows
    for index in range(len(keys) + extra):
        key = keys[index] if index < len(keys) else None
        values = grid.get(key) if key is not None else None
        cells = [(key, "locked" if row_fields else "input")]
        if row_fields:
            cells += [(row_fields[key].get("unit", ""), "locked"), (rule_text(row_fields[key]), "locked")]
        cells += [(_cell_value(values[i]) if values else None, "input") for i in range(len(columns))]
        row = out.append(*cells)
        if row_fields:
            out.validate(row_fields[key], f"{first_col}{row}:{last_col}{row}", lists)
    if not row_fields and out.row >= FIRST_ROW:
        out.validate(spec.value, f"{first_col}{FIRST_ROW}:{last_col}{out.row}", lists)
        out.validate(_key_spec(spec.rows), f"A{FIRST_ROW}:A{out.row}", lists)
    out.width(1, 22)
    for i, column in enumerate(meta + columns, start=2):
        out.width(i, 22 if column in meta else max(8, len(column) + 2))
    out.ws.freeze_panes = f"{first_col}{FIRST_ROW}"


def _cell_value(value):
    if value is None:
        return None
    if isinstance(value, float) and pd.isna(value):
        return None
    if hasattr(value, "item"):  # numpy scalar
        value = value.item()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _write_readme(out: _Sheet, catalogue: Catalogue, sheets: list[str]) -> None:
    top = max(catalogue.tables[s].level for s in sheets) if sheets else 0
    source = catalogue.path.as_posix() if catalogue.path else "the catalogue"
    lines = [
        ("MOSAICA data workbook", "title"),
        f"Generated {date.today().isoformat()} from {source} (version {catalogue.version}).",
        (f"Highest access level in this workbook: {top} ({catalogue.level_name(top)})."
         + (" Do not share it outside the people entitled to that level." if top >= 2 else ""), "bold"),
        "",
        ("How to fill it in", "bold"),
        "Yellow cells are yours to fill. Grey cells are headers, units, rules or computed values: "
        "they are locked (Review > Unprotect sheet, no password, if a row must be inserted).",
        "Each table sheet opens with its title, description, provenance and value rule (rows 1-4) "
        "and its header (row 5). Plot-like tables add a description, unit and allowed-values row "
        f"(6-8) and an example row ({EXAMPLE_KEY}, row 9, ignored when read back).",
        "Drop-down lists and bounds reject most errors as you type; the validator catches the rest "
        "(keys that do not exist, duplicates, missing rows or columns, computed values).",
        "Leave a cell empty only where the allowed row says it is optional.",
        "",
        ("Then run", "bold"),
        "python scripts/validate_data_workbook.py <this file>           -> a report of every problem",
        "python scripts/validate_data_workbook.py <this file> --export  -> the .txt tables, dated "
        "and hashed, ready for the information system",
        "",
        ("Access levels", "bold"),
    ]
    for level, spec in sorted(catalogue.levels.items()):
        lines.append(f"{level} -- {spec.get('name', '')}: {' '.join(str(spec.get('rule', '')).split())}")
    lines += ["", ("Sheets", "bold")]
    for sheet in sheets:
        spec = catalogue.tables[sheet]
        lines.append(f"{sheet} -- {spec.entry.get('title', '')} (level {spec.level}, {spec.file})")
    lines += ["", ("Colours", "bold"), ("to fill in", "input"),
              ("locked: header, unit, rule, computed", "locked"), ("example, ignored", "example")]
    for line in lines:
        out.append(line)
    out.width(1, 120)


def _write_sources(out: _Sheet, catalogue: Catalogue, sheets: list[str]) -> None:
    headers = ["sheet", "file", "title", "level", "source", "owner", "licence", "vintage",
               "resolution", "level reason", "notes"]
    out.append(*((h, "header") for h in headers))
    for sheet in sheets:
        spec = catalogue.tables[sheet]
        entry = spec.entry
        notes = [f"{name}: {' '.join(str(f['note']).split())}"
                 for name, f in {**spec.fields, **spec.row_fields}.items() if f.get("note")]
        values = [sheet, spec.file, entry.get("title", ""), spec.level, entry.get("source", ""),
                  entry.get("owner", ""), entry.get("licence", ""), str(entry.get("vintage", "")),
                  entry.get("resolution", ""), entry.get("level_reason", ""), "; ".join(notes)]
        out.append(*((v, "wrap") for v in values))
    for i, width in enumerate((26, 40, 28, 6, 40, 16, 30, 10, 14, 30, 60), start=1):
        out.width(i, width)
    out.ws.freeze_panes = "A2"


# --- reading back ------------------------------------------------------------------------------

def read_workbook(path: Path, catalogue: Catalogue) -> tuple[dict[str, pd.DataFrame], dict[str, dict]]:
    """Every table sheet of a workbook as the DataFrame catalogue.validate expects.

    Returns (frames, cells): cells[sheet] maps (row key, column) to the sheet cell, for the
    report -- pass it to validate(). Computed cells come back as their formula text; the
    validator recomputes them.
    """
    wb = load_workbook(Path(path), read_only=True)
    frames: dict[str, pd.DataFrame] = {}
    cells: dict[str, dict] = {}
    for sheet in wb.sheetnames:
        if sheet in (README, SOURCES, LISTS) or sheet not in catalogue.tables:
            continue
        spec = catalogue.tables[sheet]
        rows = list(wb[sheet].iter_rows(min_row=HEADER_ROW, values_only=True))
        reader = {"set": _read_set, "mapping": _read_mapping, "records": _read_records,
                  "matrix": _read_matrix}[spec.layout]
        frames[sheet], cells[sheet] = reader(spec, rows)
    wb.close()
    return frames, cells


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip() if not isinstance(value, str) else value


def _blank(values) -> bool:
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in values)


def _read_set(spec, rows) -> pd.DataFrame:
    data = [_text(r[0]) for r in rows[FIRST_ROW - HEADER_ROW:] if r and not _blank(r[:1])]
    return pd.DataFrame({"id": data}), {}


def _read_mapping(spec, rows) -> pd.DataFrame:
    parent, child = spec.key_names
    data = [(_text(r[0]), _text(r[1])) for r in rows[FIRST_ROW - HEADER_ROW:]
            if r and not _blank(r[:2])]
    return pd.DataFrame(data, columns=[parent, child]), {}


def _read_records(spec, rows) -> pd.DataFrame:
    header = [_text(v) for v in rows[0]]
    names = header[1:]
    while names and not names[-1]:
        names.pop()
    keys, data, cells = [], [], {}
    for offset, row in enumerate(rows[RECORDS_FIRST_ROW - HEADER_ROW:]):
        values = list(row[: len(names) + 1]) + [None] * (len(names) + 1 - len(row))
        if _blank([values[0]] + [v for v, n in zip(values[1:], names)
                                 if not _is_generated(spec, n, v)]):
            continue
        excel_row = RECORDS_FIRST_ROW + offset
        key = _text(values[0])
        keys.append(key)
        data.append(values[1:])
        for i, name in enumerate(names, start=2):
            cells[(key, name)] = f"{get_column_letter(i)}{excel_row}"
    frame = pd.DataFrame(data, index=pd.Index(keys, name=spec.rows.get("name")), columns=names)
    return frame, cells


def _is_generated(spec, name, value) -> bool:
    """A computed cell of an empty spare row is a formula, not something the user typed."""
    return name in spec.fields and "computed" in spec.fields[name] and (
        value is None or (isinstance(value, str) and value.startswith("=")))


def _read_matrix(spec, rows) -> pd.DataFrame:
    header = [_text(v) for v in rows[0]]
    offset = 1
    if spec.row_fields and tuple(header[1:3]) == MATRIX_META_COLUMNS:
        offset = 3
    columns = header[offset:]
    while columns and not columns[-1]:
        columns.pop()
    keys, data, cells = [], [], {}
    for index, row in enumerate(rows[FIRST_ROW - HEADER_ROW:]):
        values = list(row) + [None] * (offset + len(columns) - len(row))
        cells_values = values[offset: offset + len(columns)]
        if _blank([values[0]] + cells_values):
            continue
        key = _text(values[0])
        keys.append(key)
        data.append(cells_values)
        for i, column in enumerate(columns):
            cells[(key, column)] = f"{get_column_letter(offset + 1 + i)}{FIRST_ROW + index}"
    frame = pd.DataFrame(data, index=pd.Index(keys, name=spec.rows.get("name")), columns=columns)
    return frame, cells
